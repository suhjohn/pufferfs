"""Native spreadsheet rows to bounded searchable text; no inference or network IO."""

from __future__ import annotations

import csv
import json
from contextlib import closing
from itertools import chain, groupby, repeat
from pathlib import Path

from extraction import CHUNK_BYTES, chunk_record, text_chunks


def column_name(number: int) -> str:
    result = ""
    while number:
        number, digit = divmod(number - 1, 26)
        result = chr(65 + digit) + result
    return result


def row_text(row_number: int, values) -> str:
    # JSON strings distinguish embedded tabs/newlines, empty cells and literal
    # delimiters. Cell addresses remain explicit even in sparse worksheets.
    return "".join(
        f"{column_name(column)}{row_number}={json.dumps(str(value), ensure_ascii=False)}\n"
        for column, value in enumerate(values, 1) if value is not None and value != ""
    )


def sheet_chunks(sheet: str, rows, *, target_bytes: int = CHUNK_BYTES):
    """Group rows, repeating a bounded first-row context, split oversized rows.

    The first nonempty row is context, not an inferred schema. Every nonempty
    cell is retained, including oversized headers; only repeated context is
    shortened. Row and column addresses are one-based.
    """
    if target_bytes < 128:
        raise ValueError("spreadsheet chunk budget must be at least 128 bytes")
    header = ""
    pending = ""
    first = last = 0
    ordinal = 0
    for row_number, values in rows:
        content = row_text(row_number, values)
        if not content:
            continue
        if not header:
            context = next(text_chunks([content.encode()], target_bytes=target_bytes // 4))["content"]
            header = "First populated row (context):\n" + context + "\nCells:\n"
        budget = target_bytes - len(header.encode())
        if pending and len((pending + content).encode()) > budget:
            yield chunk_record(header + pending, {"sheet": sheet, "row_start": first, "row_end": last}, ordinal)
            ordinal += 1
            pending = ""
        if len(content.encode()) > budget:
            for part, chunk in enumerate(text_chunks([content.encode()], target_bytes=budget)):
                yield chunk_record(header + chunk["content"], {
                    "sheet": sheet, "row_start": row_number, "row_end": row_number, "row_part": part,
                }, ordinal)
                ordinal += 1
        else:
            if not pending:
                first = row_number
            last = row_number
            pending += content
    if pending:
        yield chunk_record(header + pending, {"sheet": sheet, "row_start": first, "row_end": last}, ordinal)


def spreadsheet_chunks(path: str, *, target_bytes: int = CHUNK_BYTES):
    suffix = Path(path).suffix.lower()
    ordinal = 0

    def numbered(chunks):
        nonlocal ordinal
        for chunk in chunks:
            chunk["chunk_index"] = ordinal
            ordinal += 1
            yield chunk

    if suffix in {".csv", ".tsv"}:
        # newline='' delegates multiline quoted records to the CSV parser.
        # The parser's field-size limit intentionally rejects pathological cells
        # rather than permitting an unbounded allocation in a shared worker.
        with open(path, encoding="utf-8-sig", newline="") as source:
            rows = csv.reader(source, delimiter="\t" if suffix == ".tsv" else ",", strict=True)
            yield from numbered(sheet_chunks("Sheet1", enumerate(rows, 1), target_bytes=target_bytes))
        return
    if suffix == ".fods":
        from fods_extract import fods_rows

        with closing(fods_rows(path)) as source_rows:
            for (_, name), rows in groupby(source_rows, key=lambda row: row[0]):
                yield from numbered(sheet_chunks(name, ((row, values) for _, row, values in rows),
                                                 target_bytes=target_bytes))
        return
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        from openpyxl import load_workbook

        # Formula expressions are indexed, never evaluated. No external links
        # or macros are loaded/executed. Read-only worksheets stream rows.
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook:
                # Some exporters report incorrect dimensions; derive them from
                # the actual rows instead of dropping cells outside that range.
                sheet.reset_dimensions()
                yield from numbered(sheet_chunks(sheet.title, enumerate(sheet.values, 1), target_bytes=target_bytes))
        finally:
            workbook.close()
        return
    if suffix in {".xls", ".xlt", ".xlsb", ".ods", ".ots"}:
        from python_calamine import CalamineWorkbook

        # Detect the container from bytes, including template extensions.
        # Calamine exposes cached values, not formula expressions. It never
        # evaluates formulas or runs macros. Its native sheet is materialized;
        # iter_rows avoids a second, whole-sheet Python list, not that allocation.
        with open(path, "rb") as source, CalamineWorkbook.from_filelike(source) as workbook:
            for name in workbook.sheet_names:
                sheet = workbook.get_sheet_by_name(name)
                if sheet.start is None:
                    continue
                # iter_rows includes leading empty rows but omits the columns
                # before sheet.start. Restore those columns for cell addresses.
                first_column = sheet.start[1]
                rows = ((number, chain(repeat(None, first_column), values))
                        for number, values in enumerate(sheet.iter_rows(), 1))
                yield from numbered(sheet_chunks(name, rows, target_bytes=target_bytes))
        return
    raise ValueError(f"native spreadsheet decoder not yet implemented for {suffix}")
