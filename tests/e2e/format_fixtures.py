"""Synthetic non-media format variants; no application extraction imports."""

import json
from pathlib import Path
import struct
import subprocess
import tempfile
from lxml import etree as ET
from zipfile import ZIP_DEFLATED, ZipFile


def office_export(source, suffix, filter_name, directory):
    # Use LibreOffice's published export filters to generate actual legacy/ODF
    # encodings. Do not relabel ZIP files as binary .doc/.ppt/.xls containers.
    with tempfile.TemporaryDirectory(prefix="office-fixture-") as temporary:
        profile = (Path(temporary) / "profile").as_uri()
        subprocess.run(["soffice", f"-env:UserInstallation={profile}", "--headless",
            "--nologo", "--nodefault", "--norestore", "--convert-to", f"{suffix}:{filter_name}",
            "--outdir", str(directory), str(source)], check=True, timeout=90,
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    result = directory / (source.stem + "." + suffix)
    assert result.is_file() and result.stat().st_size, f"fixture export failed: {suffix}"
    return result


def package_variant(source, target, part, content_type):
    # Macro-capable and template variants have the same body XML but distinct
    # package content types. These valid fixtures contain no active VBA project.
    # Preserve the original namespace declarations like the Office writers;
    # LibreOffice's package detector rejects a prefixed Types root.
    with ZipFile(source) as original, ZipFile(target, "w", compression=ZIP_DEFLATED) as output:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "[Content_Types].xml":
                document = ET.fromstring(data)
                matches = [element for element in document if element.get("PartName") == part]
                assert len(matches) == 1
                matches[0].set("ContentType", content_type)
                data = ET.tostring(document, encoding="utf-8", xml_declaration=True)
            output.writestr(info, data)


def message_fixture(path, *, unicode):
    # MS-OXMSG sections 2.4.1.1 and 2.4.2: 32-byte message header, then
    # 16-byte tagged properties and individual string streams. OleWriter is
    # solely a compound-file writer here, not the application's message parser.
    from extract_msg.ole_writer import OleWriter

    writer = OleWriter()
    kind, encoding, terminator = (0x001F, "utf-16le", 2) if unicode else (0x001E, "cp1252", 1)
    properties = [struct.pack("<IIII", 0x340D0003, 6, 0x40000 if unicode else 0, 0),
                  struct.pack("<IIII", 0x3FFD0003, 6, 1252, 0)]
    values = {0x001A: "IPM.Note", 0x0037: "Orchid correspondence café",
        0x0C1A: "Synthetic Sender", 0x0C1F: "sender@example.invalid",
        0x0E04: "reader@example.invalid", 0x1000: "Orchid observatory telescope correspondence. Café notes.",
        0x007D: "From: sender@example.invalid\r\nTo: reader@example.invalid\r\nSubject: Orchid correspondence\r\n"}
    for property_id, value in values.items():
        tag = (property_id << 16) | kind
        data = value.encode(encoding)
        writer.addEntry(f"__substg1.0_{tag:08X}", data)
        properties.append(struct.pack("<IIII", tag, 6, len(data) + terminator, 0))
    writer.addEntry("__properties_version1.0", bytes(32) + b"".join(properties))
    writer.write(str(path))


def create(directory):
    from docx import Document
    from pptx import Presentation
    from PIL import Image, ImageDraw, ImageFont
    from pillow_heif import register_heif_opener
    import openpyxl

    root = Path(directory)
    root.mkdir(parents=True)
    expected = {}
    document = Document()
    document.add_heading("Orchid observatory", 0)
    document.add_paragraph("Telescope maintenance requires careful inspection.")
    document.save(root / "word.docx")
    for suffix, filter_name in {"doc": "MS Word 97", "dot": "MS Word 97 Vorlage",
        "rtf": "Rich Text Format", "odt": "writer8", "ott": "writer8_template",
        "fodt": "OpenDocument Text Flat XML"}.items():
        office_export(root / "word.docx", suffix, filter_name, root)
    for suffix, content_type in {
        "docm": "application/vnd.ms-word.document.macroEnabled.main+xml",
        "dotx": "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
        "dotm": "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
    }.items():
        package_variant(root / "word.docx", root / f"word.{suffix}", "/word/document.xml", content_type)
    expected.update({f"word.{suffix}": {"terms": ["orchid", "telescope"], "anchor": "page_number", "anchors": [0]}
        for suffix in ("doc", "docx", "docm", "dot", "dotx", "dotm", "rtf", "odt", "ott", "fodt")})

    slides = Presentation()
    for heading in ("Orchid observatory", "Coral climate station"):
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = heading
        slide.placeholders[1].text = "Telescope maintenance and rainfall records."
    slides.save(root / "slides.pptx")
    for suffix, filter_name in {"ppt": "MS PowerPoint 97", "pps": "MS PowerPoint 97 AutoPlay",
        "pot": "MS PowerPoint 97 Vorlage", "odp": "impress8", "otp": "impress8_template",
        "fodp": "OpenDocument Presentation Flat XML"}.items():
        office_export(root / "slides.pptx", suffix, filter_name, root)
    for suffix, content_type in {
        "pptm": "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
        "ppsx": "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml",
        "ppsm": "application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml",
        "potx": "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
        "potm": "application/vnd.ms-powerpoint.template.macroEnabled.main+xml",
    }.items():
        package_variant(root / "slides.pptx", root / f"slides.{suffix}", "/ppt/presentation.xml", content_type)
    expected.update({f"slides.{suffix}": {"terms": ["orchid", "coral", "rainfall"], "anchor": "page_number", "anchors": [0, 1]}
        for suffix in ("ppt", "pptx", "pptm", "pps", "ppsx", "ppsm", "pot", "potx", "potm", "odp", "otp", "fodp")})

    cells = {"Sparse": {"D5": "Orchid sparse telescope", "E6": 42.5, "G10": "Café 🌺"},
             "Second": {"A1": "Coral rainfall station"}}
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, entries in cells.items():
        sheet = book.create_sheet(name)
        for address, value in entries.items():
            sheet[address] = value
    book.save(root / "cells.xlsx")
    for suffix, filter_name in {"xls": "MS Excel 97", "xlt": "MS Excel 97 Vorlage/Template",
        "ods": "calc8", "ots": "calc8_template", "fods": "OpenDocument Spreadsheet Flat XML"}.items():
        office_export(root / "cells.xlsx", suffix, filter_name, root)
    book.template = True
    book.save(root / "cells.xltx")
    for suffix, content_type in {"xlsm": "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
        "xltm": "application/vnd.ms-excel.template.macroEnabled.main+xml"}.items():
        package_variant(root / "cells.xlsx", root / f"cells.{suffix}", "/xl/workbook.xml", content_type)
    subprocess.run(["node", "/e2e/xlsb_fixture.cjs"], input=json.dumps({"path": str(root / "cells.xlsb"),
        "range": "A1:G10", "sheets": cells}), text=True, check=True, timeout=30)
    with ZipFile(root / "cells.xlsb") as binary:
        assert "xl/workbook.bin" in binary.namelist() and "xl/worksheets/sheet1.bin" in binary.namelist()
    serialized = {name: "".join(f"{address}={json.dumps(str(value), ensure_ascii=False)}\n"
        for address, value in entries.items()) for name, entries in cells.items()}
    expected.update({f"cells.{suffix}": {"terms": ["orchid", "coral", "café"], "sheets": serialized}
        for suffix in ("xls", "xlsx", "xlsm", "xlsb", "xlt", "xltx", "xltm", "ods", "ots", "fods")})

    register_heif_opener()
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 48)
    frames = []
    for text in ("Orchid observatory telescope", "Coral climate station rainfall"):
        picture = Image.new("RGB", (1200, 250), "white")
        ImageDraw.Draw(picture).text((20, 85), text, font=font, fill="black")
        frames.append(picture)
    single = {"png": "PNG", "jpg": "JPEG", "jpeg": "JPEG", "jfif": "JPEG", "webp": "WEBP",
        "bmp": "BMP", "heic": "HEIF", "heif": "HEIF", "avif": "AVIF", "jp2": "JPEG2000"}
    for suffix, encoding in single.items():
        frames[0].save(root / f"image.{suffix}", format=encoding)
        expected[f"image.{suffix}"] = {"terms": ["orchid", "telescope"], "anchor": "frame_number", "anchors": [0]}
    frames[0].save(root / "image.j2k", format="JPEG2000", no_jp2=True)
    expected["image.j2k"] = {"terms": ["orchid", "telescope"], "anchor": "frame_number", "anchors": [0]}
    # JPEG 2000 Part 2-compatible JPX header: change only the declared file
    # brand/compatibility list, retaining the ordinary single-image codestream.
    data = (root / "image.jp2").read_bytes()
    assert data[16:20] == b"ftyp" and data[20:24] == b"jp2 "
    (root / "image.jpx").write_bytes(data[:20] + b"jpx " + data[24:])
    expected["image.jpx"] = {"terms": ["orchid", "telescope"], "anchor": "frame_number", "anchors": [0]}
    for suffix, encoding in {"gif": "GIF", "apng": "PNG", "tif": "TIFF", "tiff": "TIFF"}.items():
        frames[0].save(root / f"image.{suffix}", format=encoding, save_all=True, append_images=frames[1:], duration=1000, loop=0)
        expected[f"image.{suffix}"] = {"terms": ["orchid", "coral", "telescope", "rainfall"], "anchor": "frame_number", "anchors": [0, 1]}
    (root / "image.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="250">'
        '<rect width="1200" height="250" fill="white"/>'
        '<text x="20" y="130" font-size="48" fill="black">Orchid observatory telescope</text></svg>')
    expected["image.svg"] = {"terms": ["orchid", "telescope"], "anchor": "frame_number", "anchors": [0]}

    for variant, unicode in (("unicode", True), ("ansi", False)):
        path = f"message-{variant}.msg"
        message_fixture(root / path, unicode=unicode)
        expected[path] = {"terms": ["orchid", "telescope", "café", "sender@example.invalid"],
                          "anchor": "record_number", "anchors": [0]}
    assert set(expected) == {path.name for path in root.iterdir() if path.is_file()}
    return expected
