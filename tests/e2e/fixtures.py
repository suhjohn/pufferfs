"""Synthetic on-disk inputs only. Never imports the application's extractors."""

from pathlib import Path
import subprocess


def create_native(directory):
    """Exercise real native decoders before any provider-backed file is queued."""
    import csv
    import json
    import openpyxl
    import xlwt
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    root = Path(directory)
    (root / "sessions").mkdir(parents=True)
    records = b''.join((json.dumps({"record": i, "nested": {"text": "Orchid 🌺 café"}},
                        ensure_ascii=False, separators=(", ", " : ")) + "\r\n").encode()
                       for i in range(1000))
    # One oversized, multibyte record plus an unfinished final record. Neither
    # ordinary nor session-named JSONL may be projected or reserialized.
    records += ('{"payload":"' + "🌺" * (1 << 18) + '"}\r\n{"partial":').encode()
    for path in ("records.jsonl", "sessions/rollout.jsonl"):
        (root / path).write_bytes(records)
    (root / "unicode.txt").write_bytes(("Orchid\r\n" + "é🌺" * 4000 + "\r\nlast line").encode())
    (root / "empty.txt").write_bytes(b"")

    long_cell = "Orchid " + "é🌺" * 1600
    expected = {}
    for suffix, delimiter in (("csv", ","), ("tsv", "\t")):
        with (root / f"cells.{suffix}").open("w", encoding="utf-8-sig", newline="") as output:
            writer = csv.writer(output, delimiter=delimiter)
            writer.writerows([["Topic", "", "Count"], ["Orchid, quoted\nsecond line\ttab", "", "42"],
                              ["", "", long_cell]])
        expected[f"cells.{suffix}"] = {"Sheet1":
            'A1="Topic"\nC1="Count"\nA2="Orchid, quoted\\nsecond line\\ttab"\nC2="42"\n'
            + "C3=" + json.dumps(long_cell, ensure_ascii=False) + "\n"}

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sparse"
    sheet["D5"] = "Orchid sparse cell"
    sheet["E6"] = "=1+2"
    sheet["G10"] = long_cell
    book.create_sheet("Second")["A1"] = "Orchid second sheet"
    book.save(root / "cells.xlsx")
    expected["cells.xlsx"] = {"Sparse": 'D5="Orchid sparse cell"\nE6="=1+2"\nG10='
                             + json.dumps(long_cell, ensure_ascii=False) + "\n",
                             "Second": 'A1="Orchid second sheet"\n'}

    legacy = xlwt.Workbook()
    legacy.add_sheet("Sparse").write(4, 3, "Orchid sparse cell")
    legacy.save(str(root / "cells.xls"))
    expected["cells.xls"] = {"Sparse": 'D5="Orchid sparse cell"\n'}

    ods = OpenDocumentSpreadsheet()
    table = Table(name="Sparse")
    table.addElement(TableRow(numberrowsrepeated=4))
    row = TableRow()
    row.addElement(TableCell(numbercolumnsrepeated=3))
    cell = TableCell(valuetype="string")
    cell.addElement(P(text="Orchid sparse cell"))
    row.addElement(cell)
    table.addElement(row)
    ods.spreadsheet.addElement(table)
    ods.save(str(root / "cells.ods"))
    expected["cells.ods"] = {"Sparse": 'D5="Orchid sparse cell"\n'}
    (root / "cells.fods").write_text('''<?xml version="1.0" encoding="UTF-8"?>
<office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
 <office:body><office:spreadsheet><table:table table:name="Sparse">
 <table:table-row table:number-rows-repeated="4"/>
 <table:table-row><table:table-cell table:number-columns-repeated="3"/>
 <table:table-cell office:value-type="string"><text:p>Orchid<text:s text:c="2"/>sparse cell</text:p></table:table-cell>
 </table:table-row></table:table></office:spreadsheet></office:body>
</office:document>''')
    expected["cells.fods"] = {"Sparse": 'D5="Orchid  sparse cell"\n'}
    (root / "mail.eml").write_text("From: sender@example.invalid\nTo: reader@example.invalid\n"
                                  "Subject: Orchid native email\n\nOrchid body.\n")
    (root / "contact.vcf").write_text("BEGIN:VCARD\nVERSION:3.0\nFN:Orchid Contact\nEND:VCARD\n")
    (root / "event.ics").write_text("BEGIN:VCALENDAR\nVERSION:2.0\nBEGIN:VEVENT\n"
                                   "SUMMARY:Orchid event\nEND:VEVENT\nEND:VCALENDAR\n")
    return sorted(str(path.relative_to(root)) for path in root.rglob("*") if path.is_file()), expected


def create(directory):
    from docx import Document
    from pptx import Presentation
    from PIL import Image, ImageDraw, ImageFont
    import openpyxl
    import pymupdf
    import xlwt

    root = Path(directory)
    root.mkdir(parents=True)
    for d in range(100):
        folder = root / f"directory-{d:03}"
        folder.mkdir()
        for f in range(10):
            (folder / f"file-{f:02}.txt").write_text(
                f"Synthetic pufferfs directory {d} file {f}.\nOrchid observatory maintenance.\n")
    (root / "append.jsonl").write_text(''.join(
        '{"record":%d,"text":"Orchid telemetry reading %d"}\n' % (i, i) for i in range(2000)))
    (root / "rewrite.txt").write_text("Original cobalt description.\n")
    (root / "remove.txt").write_text("Obsolete saffron schedule.\n")
    (root / "move.txt").write_text("Moving cedar notes.\n")
    (root / "empty.txt").write_bytes(b"")
    (root / "sample.py").write_text('print("Orchid observatory")\n')
    (root / "sample.eml").write_text("From: sender@example.invalid\nTo: reader@example.invalid\n"
                                     "Subject: Orchid observatory\n\nAnnual maintenance.\n")
    (root / "sample.vcf").write_text("BEGIN:VCARD\nVERSION:3.0\nFN:Orchid Observatory\nEND:VCARD\n")
    (root / "sample.ics").write_text("BEGIN:VCALENDAR\nVERSION:2.0\nBEGIN:VEVENT\n"
                                     "SUMMARY:Orchid observatory\nEND:VEVENT\nEND:VCALENDAR\n")
    (root / ".env").write_text("SYNTHETIC_SECRET=must-never-be-captured\n")
    # Cross the actual 32 MiB capture pack threshold, not a test-only override.
    with (root / "large.txt").open("wb") as output:
        line = b"Orchid observatory large text line.\n"
        for _ in range((33 << 20) // (len(line) * 1024) + 1):
            output.write(line * 1024)

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Observations"
    sheet.append(["Location", "Count"])
    sheet.append(["Orchid observatory", 42])
    book.save(root / "sample.xlsx")
    legacy = xlwt.Workbook()
    sheet = legacy.add_sheet("Observations")
    sheet.write(0, 0, "Orchid observatory")
    sheet.write(0, 1, 42)
    legacy.save(str(root / "sample.xls"))
    for extension, delimiter in [("csv", ","), ("tsv", "\t")]:
        (root / f"sample.{extension}").write_text(f"Location{delimiter}Count\nOrchid observatory{delimiter}42\n")

    doc = Document()
    doc.add_heading("Orchid observatory", 0)
    doc.add_paragraph("Annual maintenance requires inspection of the telescope.")
    doc.save(root / "sample.docx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Orchid observatory"
    slide.placeholders[1].text = "Annual telescope maintenance."
    presentation.save(root / "sample.pptx")
    with pymupdf.open() as pdf:
        for page in range(2):
            pdf.new_page().insert_text((72, 100), f"Orchid observatory page {page + 1}", fontsize=24)
        pdf.save(root / "sample.pdf")
    picture = Image.new("RGB", (1200, 300), "white")
    ImageDraw.Draw(picture).text((30, 80), "Orchid observatory", fill="black",
        font=ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 56))
    for extension in ("png", "jpg", "webp", "tiff"):
        picture.save(root / f"sample.{extension}")

    expectations = {f"sample.{suffix}": {"terms": ["orchid"]} for suffix in (
        "py", "eml", "vcf", "ics", "xlsx", "xls", "csv", "tsv", "docx", "pptx",
        "pdf", "png", "jpg", "webp", "tiff")}
    expectations.update(create_media(root))
    files = sorted(str(path.relative_to(root)) for path in root.rglob("*") if path.is_file() and path.name != ".env")
    return files, expectations


def create_media(directory):
    """Identical audio fixtures for the corpus and focused provider diagnostics."""
    root = Path(directory)
    root.mkdir(parents=True, exist_ok=True)
    # eSpeak is solely a fixture generator; production decoding uses FFmpeg.
    subprocess.run(["espeak-ng", "-v", "en-us", "-s", "130", "-w", str(root / "sample.wav"),
                    "Welcome to the Orchid observatory. The telescope is ready."], check=True)
    subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-i", str(root / "sample.wav"),
                    str(root / "sample.mp3")], check=True)
    subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-f", "lavfi", "-i",
                    "color=c=black:s=320x240:r=10", "-i", str(root / "sample.wav"),
                    "-shortest", "-c:v", "libx264", "-c:a", "aac", str(root / "sample.mp4")], check=True)
    return {f"sample.{suffix}": {"terms": ["observatory", "telescope"], "diarized": True,
                                "query": "telescope", "clip_count": 1} for suffix in ("wav", "mp3", "mp4")}


def create_extended_media(directory):
    """Actual alternate containers/codecs plus speech across multiple minute clips."""
    import tempfile
    import wave
    from amr_fixture import create_amr

    root = Path(directory)
    expected = create_media(root)
    # Explicit fixture encodings, not extension-only renames or an application
    # format table imported into the tests. The runner's FFmpeg has an AMR
    # decoder but no encoder; a fixture-only OpenCORE encoder supplies its input.
    audio_encodings = {
        "m4a": ["-c:a", "aac"], "m4b": ["-c:a", "aac", "-f", "ipod"],
        "aac": ["-c:a", "aac"], "flac": ["-c:a", "flac"],
        "ogg": ["-c:a", "libvorbis"], "oga": ["-c:a", "flac", "-f", "ogg"],
        "opus": ["-c:a", "libopus"], "aif": ["-c:a", "pcm_s16be"],
        "aiff": ["-c:a", "pcm_s16be"], "wma": ["-c:a", "wmav2"],
    }
    for suffix, encoding in audio_encodings.items():
        subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-i", str(root / "sample.wav"),
                        *encoding, str(root / ("sample." + suffix))], check=True)
    video_encodings = {
        "mov": ["-c:v", "libx264", "-c:a", "aac"],
        "m4v": ["-c:v", "libx264", "-c:a", "aac", "-f", "mp4"],
        "mkv": ["-c:v", "libx264", "-c:a", "flac"],
        "webm": ["-c:v", "libvpx-vp9", "-c:a", "libopus"],
        "avi": ["-c:v", "mpeg4", "-c:a", "pcm_s16le"],
        "mpeg": ["-c:v", "mpeg2video", "-c:a", "mp2", "-ar", "44100", "-f", "mpeg"],
        "mpg": ["-c:v", "mpeg2video", "-c:a", "mp2", "-ar", "44100", "-f", "mpeg"],
        "wmv": ["-c:v", "wmv2", "-c:a", "wmav2"],
        "flv": ["-c:v", "flv", "-c:a", "libmp3lame", "-ar", "44100"],
        "3gp": ["-c:v", "mpeg4", "-c:a", "aac", "-ar", "16000"],
        "mts": ["-c:v", "libx264", "-c:a", "aac", "-f", "mpegts"],
        "m2ts": ["-c:v", "libx264", "-c:a", "aac", "-f", "mpegts"],
        "mxf": ["-vf", "scale=720:576", "-c:v", "mpeg2video", "-pix_fmt", "yuv422p",
                "-b:v", "50000k", "-c:a", "pcm_s16le", "-ar", "48000", "-f", "mxf"],
    }
    for suffix, encoding in video_encodings.items():
        subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-f", "lavfi", "-i",
                        "color=c=black:s=320x240:r=25", "-i", str(root / "sample.wav"),
                        "-shortest", *encoding, str(root / ("sample." + suffix))], check=True)
    for suffix in (*audio_encodings, *video_encodings):
        expected["sample." + suffix] = {"terms": ["observatory", "telescope"],
            "diarized": True, "query": "telescope", "clip_count": 1}
    create_amr(root / "sample.wav", root / "sample.amr")
    expected["sample.amr"] = {"terms": ["observatory", "telescope"],
        "diarized": True, "query": "telescope", "clip_count": 1}

    with tempfile.TemporaryDirectory(prefix="pufferfs-media-fixture-") as temporary:
        voice = Path(temporary) / "second-voice.wav"
        subprocess.run(["espeak-ng", "-v", "en-gb+f3", "-s", "135", "-w", str(voice),
                        "The coral climate station is recording rainfall."], check=True)
        # Quiet intervals make the timestamp offset unambiguous without
        # requiring Gemini to repeat hundreds of identical spoken sentences.
        timeline = ("[0:a]aresample=16000,asplit=2[a0][a1];[a1]adelay=285000[a2];"
                    "[1:a]aresample=16000,asplit=2[b0][b1];[b0]adelay=15000[b2];"
                    "[b1]adelay=305000[b3];[a0][a2][b2][b3]amix=inputs=4:normalize=0,"
                    "apad=whole_len=5040000,atrim=end_sample=5040000,asetpts=N/SR/TB[out]")
        subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-i", str(root / "sample.wav"),
                        "-i", str(voice), "-filter_complex", timeline, "-map", "[out]",
                        "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le", "-fs", str(5040000 * 2 + 4096),
                        str(root / "long.wav")], check=True, timeout=60)
    # Bound the fixture by sample count, independently of filter timestamps.
    # Unlimited apad with time-based atrim produced hours of silence on this
    # FFmpeg build. Verify the actual header/size before any capture begins.
    with wave.open(str(root / "long.wav"), "rb") as recording:
        assert (recording.getframerate(), recording.getnchannels(), recording.getsampwidth(), recording.getnframes()) == (16000, 1, 2, 5040000)
    assert (root / "long.wav").stat().st_size <= 5040000 * 2 + 4096
    expected["long.wav"] = {"terms": ["observatory", "telescope", "climate", "rainfall"],
        "diarized": True, "query": "rainfall", "clip_count": 6, "clips": [
            {"start_seconds": 0, "end_seconds": 60, "terms": ["observatory", "telescope", "climate"]},
            {"start_seconds": 60, "end_seconds": 120, "terms": [], "silent": True},
            {"start_seconds": 120, "end_seconds": 180, "terms": [], "silent": True},
            {"start_seconds": 180, "end_seconds": 240, "terms": [], "silent": True},
            {"start_seconds": 240, "end_seconds": 300, "terms": ["observatory", "telescope"]},
            {"start_seconds": 300, "end_seconds": 315, "terms": ["climate", "rainfall"]},
        ], "utterances": [
            {"terms": ["observatory"], "start_seconds": 0, "tolerance_seconds": 3},
            {"terms": ["climate"], "start_seconds": 15, "tolerance_seconds": 3},
            {"terms": ["observatory"], "start_seconds": 285, "tolerance_seconds": 3},
            {"terms": ["climate"], "start_seconds": 305, "tolerance_seconds": 3},
        ]}
    return expected
